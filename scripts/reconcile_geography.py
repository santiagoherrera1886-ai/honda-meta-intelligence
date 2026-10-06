#!/usr/bin/env python3
"""Validate daily spend against an Excel export and retain typed Meta geography.
Usage: python scripts/reconcile_geography.py workbook.xlsx /path/to/CO.zip
No ad performance or original snapshot values are modified.
"""
import sys, json, gzip, base64, re, hashlib, zipfile, unicodedata
from pathlib import Path
from collections import defaultdict, Counter
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
KINDS=('cities','regions','countries','subcities','neighborhoods','places','medium_geo_areas','custom_locations','zips')
ACCENTS={'Narino':'Nariño','Atlantico':'Atlántico','Bolivar':'Bolívar','Cordoba':'Córdoba','Caqueta':'Caquetá','Choco':'Chocó','Boyaca':'Boyacá','Quindio':'Quindío','Guainia':'Guainía'}
def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFD',str(value or '').strip().lower()) if unicodedata.category(c)!='Mn')
def clean(value):return str(value or '').strip()
def rows(sheet):
    it=sheet.iter_rows(values_only=True);headers=next(it)
    return [dict(zip(headers,row)) for row in it if any(x is not None for x in row)]
def geo_category(g):
    types={k for k in KINDS if g.get(k)}
    return 'city' if types=={'cities'} else 'region' if types=={'regions'} else 'country' if types=={'countries'} else 'mixed' if types else 'missing'
def main(workbook,gazetteer):
    wb=load_workbook(workbook,read_only=True,data_only=True)
    segments=rows(wb['META_SEGMENTACIONES']);perf=rows(wb['META_PERFORMANCE_DIA'])
    byid={clean(r['Ad Set ID']):r for r in segments}
    assert len(byid)==len(segments),'Duplicate Ad Set IDs in segmentation sheet'
    snapshot=json.loads(gzip.decompress(base64.b64decode(re.search(r'"(.*?)"',(ROOT/'data-packed.js').read_text())[1])))
    daily={(clean(r['Fecha'])[:10],clean(r['Ad Set ID'])):r for r in perf}
    assert len(daily)==len(perf)==len(snapshot['perf']), 'Different daily row count or duplicate day/Ad Set'
    spent=defaultdict(float)
    for row in snapshot['perf']:
        aid=snapshot['dict']['adsetId'][row[3]];source=daily[(row[0],aid)]
        assert float(source['Inversión'] or 0)==row[8],f'Different spend for {aid} {row[0]}'
        assert float(source['Leads'] or 0)==row[16],f'Different leads for {aid} {row[0]}'
        spent[aid]+=row[8]
    gazetteer_rows=[line.split('\t') for line in zipfile.ZipFile(gazetteer).read('CO.txt').decode().splitlines()]
    region_codes={}
    for row in gazetteer_rows:
        if row[7]=='ADM1':
            canonical=re.sub(r'^Departamento (?:del |de )','',row[1])
            region_codes[norm(canonical)]=row[10]
            if 'bogota' in norm(row[1]):region_codes['distrito especial']=row[10]
    place_index=defaultdict(dict);alias_index=defaultdict(dict)
    for row in gazetteer_rows:
        if row[7] not in ('PPLC','PPLA','PPLA2'):continue
        point={'lat':float(row[4]),'lon':float(row[5]),'geonamesId':row[0]}
        for name in (row[1],row[2]):place_index[(norm(name),row[10])][row[0]]=point
        for name in row[3].split(','):alias_index[(norm(name),row[10])][row[0]]=point
    locations={};adsets={};problems=[]
    def register(kind,item):
        obj={'key':item,'name':'Colombia' if item=='CO' else item,'country':item} if isinstance(item,str) else item
        identifier=clean(obj.get('key'))
        if not identifier:
            identifier=hashlib.sha256(json.dumps(obj,sort_keys=True).encode()).hexdigest()[:16]
        ref=kind+':'+identifier
        region=ACCENTS.get(clean(obj.get('region')),clean(obj.get('region')))
        name=ACCENTS.get(clean(obj.get('name')),clean(obj.get('name'))) or identifier
        meta={'id':identifier,'kind':kind,'name':name,'region':region,'regionId':clean(obj.get('region_id')),'country':clean(obj.get('country'))}
        meta['label']=name+(' · '+region if kind=='cities' and region else '')
        if kind=='cities':
            admin=region_codes.get(norm(region))
            candidates=place_index[(norm(name),admin)] or alias_index[(norm(name),admin)]
            if len(candidates)==1:
                point=next(iter(candidates.values()))
                if point['lon']>-80:meta['point']=point
        previous=locations.get(ref)
        if previous and any(norm(previous.get(k))!=norm(meta.get(k)) for k in ('name','region','kind')):
            problems.append((ref,previous,meta))
        locations[ref]=meta
        return ref
    for aid in sorted(spent):
        if aid not in byid:continue
        source=byid[aid];targeting=json.loads(source['Targeting completo JSON'] or '{}')
        included=targeting.get('geo_locations',{});excluded=targeting.get('excluded_geo_locations',{})
        unknown=set(included)-set(KINDS)-{'location_types'}
        assert not unknown,f'Unhandled geography types: {unknown}'
        refs=[];excluded_refs=[];radii={};counts={}
        for kind in KINDS:
            values=included.get(kind,[])
            if values:counts[kind]=len(values)
            for item in values:
                ref=register(kind,item);refs.append(ref)
                if isinstance(item,dict) and item.get('radius') is not None:radii[ref]={'radius':item['radius'],'unit':item.get('distance_unit','')}
            for item in excluded.get(kind,[]):excluded_refs.append(register(kind,item))
        adsets[aid]={'locations':list(dict.fromkeys(refs)),'excluded':list(dict.fromkeys(excluded_refs)),'counts':counts,'category':geo_category(included),'radii':radii,'locationTypes':included.get('location_types',[])}
    assert not problems, f'Conflicting geographic IDs: {problems[:2]}'
    included_refs={ref for v in adsets.values() for ref in v['locations']}
    city_refs=[ref for ref in included_refs if locations[ref]['kind']=='cities']
    categories={k:{'adsets':0,'spend':0} for k in ('city','region','country','mixed','missing')}
    for aid,amount in spent.items():
        kind=adsets.get(aid,{}).get('category','missing');categories[kind]['adsets']+=1;categories[kind]['spend']+=amount
    reconciliation=[]
    for aid,sg in snapshot['seg'].items():
        if 'Nariño' not in sg.get('ci',[]):continue
        geo=adsets[aid];narino=[ref for ref in geo['locations'] if locations[ref]['kind']=='cities' and locations[ref]['name']=='Nariño']
        assert len(narino)==1
        reconciliation.append({'adsetId':aid,'adset':snapshot['dict']['adset'][sg['a']],'location':locations[narino[0]]['label'],'locationKey':narino[0],'spend':spent[aid],'oldAllocated':spent[aid]/len(sg['ci']),'counts':geo['counts']})
    assert len({r['locationKey'] for r in reconciliation})==3
    report={'source':Path(workbook).name,'sha256':hashlib.sha256(Path(workbook).read_bytes()).hexdigest(),'dailyRows':len(perf),'matchingDailyRows':len(perf),'period':[min(k[0] for k in daily),max(k[0] for k in daily)],'totalSpend':sum(spent.values()),'totalLeads':sum(float(r['Leads'] or 0) for r in perf),'performanceAdsets':len(spent),'matchedAdsets':len(adsets),'missingAdsets':sorted(set(spent)-set(adsets)),'cityKeys':len(city_refs),'geocodedCities':sum('point' in locations[ref] for ref in city_refs),'categories':categories,'narino':reconciliation,'previousNarinoAllocation':sum(r['oldAllocated'] for r in reconciliation)}
    payload={'schema':2,'source':{'file':Path(workbook).name,'sha256':report['sha256'],'verifiedDailyRows':len(perf)},'locations':dict(sorted(locations.items())),'adsets':adsets}
    (ROOT/'geo-targeting.js').write_text('// Meta geography reconciled by Ad Set ID. See docs/GEO_RECONCILIATION.md.\nconst GEO_TARGETING='+json.dumps(payload,ensure_ascii=False,separators=(',',':'))+';\n')
    (ROOT/'docs/geography-reconciliation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='narino'},ensure_ascii=False,indent=2))
    print('Nariño identities:',[(ref,locations[ref].get('point')) for ref in city_refs if locations[ref]['name']=='Nariño'])
if __name__=='__main__':main(sys.argv[1],sys.argv[2])
